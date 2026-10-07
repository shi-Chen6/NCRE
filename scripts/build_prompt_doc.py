# -*- coding: utf-8 -*-
"""把 prompts/ 下的 Markdown（提示词文档）转成排版清爽的 Word 文档。

用法：
    python build_prompt_doc.py                       # 默认转换 20 题知识点讲解提示词
    python build_prompt_doc.py <src.md> [dst.docx]   # 转换指定文件
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "prompts")
DEFAULT_SRC = os.path.join(BASE, "第1套题20题知识点讲解提示词.md")

SRC = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC
DST = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(SRC)[0] + ".docx"

FONT = "微软雅黑"
C_TITLE = RGBColor(0x1F, 0x38, 0x64)
C_H2 = RGBColor(0x1F, 0x38, 0x64)
C_H3 = RGBColor(0xC0, 0x1B, 0x1B)
C_BODY = RGBColor(0x33, 0x33, 0x33)
C_GRAY = RGBColor(0x77, 0x77, 0x77)


def style_run(run, size=10.5, bold=False, color=C_BODY, italic=False, mono=False):
    name = "Consolas" if mono else FONT
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_runs(p, text, size=10.5, color=C_BODY, bold=False):
    """处理 **加粗** 与 `代码` 内联标记。"""
    for seg in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            style_run(p.add_run(seg[2:-2]), size=size, bold=True, color=color)
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            style_run(p.add_run(seg[1:-1]), size=size - 0.5, color=RGBColor(0x0B, 0x6E, 0x6E), mono=True)
        else:
            style_run(p.add_run(seg), size=size, bold=bold, color=color)


def para(doc, after=3, before=0, line=1.35, indent=0, left=0, align=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(indent)
    if left:
        pf.left_indent = Pt(left)
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def hrule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(4)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "D0D0D0")
    bottom.set(qn("w:space"), "1")
    bdr.append(bottom)
    pPr.append(bdr)


def shade_para(p, hexcolor):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    pPr.append(el)


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = r[j] if j < len(r) else ""
            cell = cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.2
            add_runs(p, txt, size=9, color=C_TITLE if i == 0 else C_BODY, bold=(i == 0))
            if i == 0:
                shade(cell, "ECF1F8")
            elif i % 2 == 0:
                shade(cell, "F8F9FB")
    para(doc, after=6)


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    i = 0
    first_h1 = True
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if not s:
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        if s == "---":
            hrule(doc)
            i += 1
            continue

        if s.startswith("# "):
            p = para(doc, after=6, before=0 if first_h1 else 14,
                     align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
            add_runs(p, s[2:], size=19, color=C_TITLE, bold=True)
            first_h1 = False
            i += 1
            continue

        if s.startswith("## "):
            p = para(doc, after=6, before=14, keep=True)
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "10")
            bottom.set(qn("w:color"), "1F3864")
            bottom.set(qn("w:space"), "3")
            bdr.append(bottom)
            pPr.append(bdr)
            add_runs(p, s[3:], size=14.5, color=C_H2, bold=True)
            i += 1
            continue

        if s.startswith("### "):
            txt = s[4:]
            m = re.match(r"^【(提示词|生图)\s*(\d+)】(.*)$", txt)
            p = para(doc, after=5, before=13, keep=True)
            if m:
                style_run(p.add_run(f"【{m.group(1)} {m.group(2)}】"), size=12.5, bold=True, color=C_H3)
                style_run(p.add_run(m.group(3)), size=12.5, bold=True, color=C_H2)
            else:
                add_runs(p, txt, size=12.5, color=C_H2, bold=True)
            i += 1
            continue

        if s.startswith("> "):
            body = s[2:]
            if len(body) > 60:
                # 长引用块 = 提示词正文，做成浅底提示框
                p = para(doc, after=6, before=2, left=10, line=1.45)
                shade_para(p, "F2F6FB")
                add_runs(p, body, size=10)
            else:
                p = para(doc, after=2, before=0, left=12, line=1.3)
                add_runs(p, body, size=9.5, color=C_GRAY)
            i += 1
            continue

        if re.match(r"^\d+\.\s", s):
            p = para(doc, after=2, left=16, line=1.35)
            p.paragraph_format.first_line_indent = Pt(-16)
            m = re.match(r"^(\d+)\.\s+(.*)$", s)
            style_run(p.add_run(f"{m.group(1)}. "), size=10.5, bold=True, color=C_TITLE)
            add_runs(p, m.group(2), size=10.5)
            i += 1
            continue

        if s.startswith("- "):
            p = para(doc, after=2, left=14, line=1.35)
            p.paragraph_format.first_line_indent = Pt(-14)
            style_run(p.add_run("• "), size=10.5, color=C_TITLE, bold=True)
            add_runs(p, s[2:], size=10.5)
            i += 1
            continue

        p = para(doc, after=3, line=1.35)
        add_runs(p, s, size=10.5)
        i += 1

    doc.save(DST)
    print("saved:", DST, os.path.getsize(DST), "bytes")


if __name__ == "__main__":
    main()
